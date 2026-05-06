from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from datetime import datetime
from agent.models import Factura


COLOR_HEADER = "1F4E79"   # Azul oscuro
COLOR_TOTAL  = "D6E4F0"   # Azul claro
COLOR_ALT    = "F2F2F2"   # Gris alternado


def estilo_header(cell):
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor=COLOR_HEADER)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def borde_delgado():
    lado = Side(style="thin", color="CCCCCC")
    return Border(left=lado, right=lado, top=lado, bottom=lado)


def exportar_excel(facturas: list[Factura], ruta_salida: str) -> str:
    """
    Genera un Excel profesional con:
    - Hoja de detalle de facturas
    - Hoja de resumen por categoría
    - Gráfico de barras de gastos
    """
    wb = Workbook()

    # ── HOJA 1: DETALLE ──────────────────────────────────────────
    ws = wb.active
    ws.title = "Facturas"

    headers = [
        "N°", "Archivo", "Proveedor", "RUC", "Fecha",
        "Descripción", "Categoría", "Subtotal", "IGV", "Total", "Moneda"
    ]

    # Escribir headers
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        estilo_header(cell)
        cell.border = borde_delgado()

    # Escribir datos
    for i, f in enumerate(facturas, 2):
        fila = [
            i - 1, f.archivo, f.proveedor, f.ruc_proveedor, f.fecha,
            f.descripcion, f.categoria, f.subtotal, f.igv, f.total, f.moneda
        ]
        for col, valor in enumerate(fila, 1):
            cell = ws.cell(row=i, column=col, value=valor)
            cell.border = borde_delgado()
            cell.alignment = Alignment(vertical="center")
            # Formato moneda para columnas numéricas
            if col in [8, 9, 10] and valor:
                cell.number_format = '#,##0.00'
            # Filas alternadas
            if i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=COLOR_ALT)

    # Fila de totales
    fila_total = len(facturas) + 2
    ws.cell(fila_total, 1, "TOTAL").font = Font(bold=True)
    for col, campo in [(8, "subtotal"), (9, "igv"), (10, "total")]:
        total = sum(getattr(f, campo) or 0 for f in facturas)
        cell = ws.cell(fila_total, col, total)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=COLOR_TOTAL)
        cell.number_format = '#,##0.00'
        cell.border = borde_delgado()

    # Ajustar ancho de columnas
    anchos = [5, 20, 25, 15, 12, 30, 15, 12, 12, 12, 8]
    for col, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(col)].width = ancho

    ws.row_dimensions[1].height = 25
    ws.freeze_panes = "A2"  # Congelar header

    # ── HOJA 2: RESUMEN POR CATEGORÍA ────────────────────────────
    ws2 = wb.create_sheet("Resumen")

    ws2.cell(1, 1, "Categoría").font = Font(bold=True, color="FFFFFF", size=11)
    ws2.cell(1, 1).fill = PatternFill("solid", fgColor=COLOR_HEADER)
    ws2.cell(1, 2, "Total (S/)").font = Font(bold=True, color="FFFFFF", size=11)
    ws2.cell(1, 2).fill = PatternFill("solid", fgColor=COLOR_HEADER)
    ws2.cell(1, 3, "# Facturas").font = Font(bold=True, color="FFFFFF", size=11)
    ws2.cell(1, 3).fill = PatternFill("solid", fgColor=COLOR_HEADER)

    # Agrupar por categoría
    categorias: dict = {}
    for f in facturas:
        cat = f.categoria or "Sin categoría"
        if cat not in categorias:
            categorias[cat] = {"total": 0, "count": 0}
        categorias[cat]["total"] += f.total or 0
        categorias[cat]["count"] += 1

    for i, (cat, datos) in enumerate(
        sorted(categorias.items(), key=lambda x: x[1]["total"], reverse=True), 2
    ):
        ws2.cell(i, 1, cat).border = borde_delgado()
        cell_total = ws2.cell(i, 2, datos["total"])
        cell_total.number_format = '#,##0.00'
        cell_total.border = borde_delgado()
        ws2.cell(i, 3, datos["count"]).border = borde_delgado()

    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 15
    ws2.column_dimensions["C"].width = 12

    # Gráfico de barras
    chart = BarChart()
    chart.type = "col"
    chart.title = "Gasto por Categoría"
    chart.y_axis.title = "S/"
    chart.style = 10

    n_cats = len(categorias)
    data = Reference(ws2, min_col=2, min_row=1, max_row=n_cats + 1)
    cats = Reference(ws2, min_col=1, min_row=2, max_row=n_cats + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.width = 20
    chart.height = 12
    ws2.add_chart(chart, "E2")

    # ── HOJA 3: META ─────────────────────────────────────────────
    ws3 = wb.create_sheet("Info")
    ws3["A1"] = "Agente de Facturas"
    ws3["A1"].font = Font(bold=True, size=14, color=COLOR_HEADER)
    ws3["A2"] = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws3["A3"] = f"Total facturas procesadas: {len(facturas)}"
    ws3["A4"] = f"Total gasto: S/ {sum(f.total or 0 for f in facturas):,.2f}"

    # Guardar
    wb.save(ruta_salida)
    return ruta_salida